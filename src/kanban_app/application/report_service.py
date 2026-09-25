from __future__ import annotations

import calendar
import csv
import subprocess
from datetime import date, datetime
from pathlib import Path
from typing import Sequence

import openpyxl
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from kanban_app.application.dto import (
    MonthlyOpItemDTO,
    MonthlyReportSummaryDTO,
    MonthlySectorStatDTO,
    MonthlyWeekStatDTO,
    SectorDTO,
)
from kanban_app.formatting import format_br_date
from kanban_app.infrastructure.db.repositories import ProductionRepository

MONTH_NAMES = (
    "",
    "Janeiro",
    "Fevereiro",
    "Março",
    "Abril",
    "Maio",
    "Junho",
    "Julho",
    "Agosto",
    "Setembro",
    "Outubro",
    "Novembro",
    "Dezembro",
)


class MonthlyReportService:
    def __init__(self, repository: ProductionRepository):
        self.repository = repository

    def build_summary(
        self,
        year: int,
        month: int,
        *,
        reference_date: date | None = None,
    ) -> MonthlyReportSummaryDTO:
        today = reference_date or date.today()
        month = max(1, min(12, int(month)))
        year = max(2000, min(2100, int(year)))

        _, last_day = calendar.monthrange(year, month)
        periodo_inicio = date(year, month, 1)
        periodo_fim = date(year, month, last_day)

        is_current_month = (year == today.year and month == today.month)
        is_mes_fechado = not is_current_month and (periodo_fim < today)

        raw_ops, sectors = self.repository.get_monthly_report_raw(periodo_inicio, periodo_fim)

        items: list[MonthlyOpItemDTO] = []
        created_count = 0
        created_pcs = 0
        concluded_count = 0
        concluded_pcs = 0
        on_time_count = 0
        delayed_count = 0
        lead_time_days_total = 0
        lead_time_ops_count = 0
        in_production_now = 0
        in_delay_now = 0
        remaining_forecast = 0

        sector_counts: dict[str, int] = {s.nome: 0 for s in sectors}
        sector_pcs: dict[str, int] = {s.nome: 0 for s in sectors}

        seen_op_ids: set[int] = set()

        for op in raw_ops:
            op_id = int(op["id"])
            op_created_at: datetime | None = op["created_at"]
            op_completed_at: datetime | None = op["completed_at"]
            d_inicio: date | None = op["data_inicio"]
            d_entrega: date | None = op["data_entrega"]
            qtd = op["quantidade"] if op["quantidade"] is not None else 0
            sec_nome = str(op["setor_nome"] or "Sem setor")
            status_str = str(op["status"] or "")

            created_in_month = False
            if op_created_at and periodo_inicio <= op_created_at.date() <= periodo_fim:
                created_in_month = True
            elif d_inicio and periodo_inicio <= d_inicio <= periodo_fim:
                created_in_month = True

            completed_in_month = False
            if op_completed_at and periodo_inicio <= op_completed_at.date() <= periodo_fim:
                completed_in_month = True

            is_active_current = is_current_month and not op["archived"] and status_str != "CONCLUIDO"

            if not (created_in_month or completed_in_month or (is_current_month and is_active_current)):
                continue

            if op_id in seen_op_ids:
                continue
            seen_op_ids.add(op_id)

            if created_in_month:
                created_count += 1
                created_pcs += qtd

            entregue_no_prazo: bool | None = None
            esta_em_atraso = False
            dias_producao: int | None = None
            categoria = "EM_PRODUCAO"

            if completed_in_month:
                concluded_count += 1
                concluded_pcs += qtd

                if d_entrega is not None and op_completed_at is not None:
                    if op_completed_at.date() <= d_entrega:
                        entregue_no_prazo = True
                        on_time_count += 1
                        categoria = "CONCLUIDA_NO_PRAZO"
                    else:
                        entregue_no_prazo = False
                        delayed_count += 1
                        categoria = "CONCLUIDA_COM_ATRASO"
                else:
                    entregue_no_prazo = True
                    on_time_count += 1
                    categoria = "CONCLUIDA_NO_PRAZO"

                if d_inicio and op_completed_at:
                    diff = (op_completed_at.date() - d_inicio).days
                    dias_producao = max(0, diff)
                    lead_time_days_total += dias_producao
                    lead_time_ops_count += 1
            else:
                if is_current_month and is_active_current:
                    in_production_now += 1
                    if d_entrega and d_entrega < today:
                        esta_em_atraso = True
                        in_delay_now += 1
                        categoria = "EM_ATRASO"
                    else:
                        categoria = "EM_PRODUCAO"
                        if d_entrega and today <= d_entrega <= periodo_fim:
                            remaining_forecast += 1

                    if d_inicio:
                        dias_producao = max(0, (today - d_inicio).days)

            if sec_nome not in sector_counts:
                sector_counts[sec_nome] = 0
                sector_pcs[sec_nome] = 0
            sector_counts[sec_nome] += 1
            sector_pcs[sec_nome] += qtd

            items.append(
                MonthlyOpItemDTO(
                    op_id=op_id,
                    numero_op=str(op["numero_op"]),
                    cliente=str(op["cliente"]),
                    modelo=str(op["modelo"]),
                    quantidade=op["quantidade"],
                    voltagem=str(op["voltagem"] or ""),
                    setor_nome=sec_nome,
                    status=status_str,
                    data_inicio=d_inicio,
                    data_entrega=d_entrega,
                    completed_at=op_completed_at,
                    dias_producao=dias_producao,
                    entregue_no_prazo=entregue_no_prazo,
                    esta_em_atraso=esta_em_atraso,
                    categoria=categoria,
                )
            )

        items.sort(key=lambda x: (x.data_entrega or date.max, x.numero_op))

        taxa_pontualidade = 100.0
        if concluded_count > 0:
            taxa_pontualidade = round((on_time_count / concluded_count) * 100.0, 1)

        lead_time_medio = 0.0
        if lead_time_ops_count > 0:
            lead_time_medio = round(lead_time_days_total / lead_time_ops_count, 1)

        sector_stats: list[MonthlySectorStatDTO] = []
        sec_color_map = {s.nome: s.cor for s in sectors}
        sec_id_map = {s.nome: s.id for s in sectors}
        for s_nome, count in sector_counts.items():
            if count > 0:
                sector_stats.append(
                    MonthlySectorStatDTO(
                        setor_id=sec_id_map.get(s_nome),
                        setor_nome=s_nome,
                        cor=sec_color_map.get(s_nome, "#3b82f6"),
                        total_ops=count,
                        total_quantidade=sector_pcs.get(s_nome, 0),
                    )
                )
        sector_stats.sort(key=lambda s: s.total_ops, reverse=True)

        week_stats: list[MonthlyWeekStatDTO] = []
        cur_day = 1
        week_idx = 1
        while cur_day <= last_day:
            w_start = date(year, month, cur_day)
            w_end_day = min(cur_day + 6, last_day)
            w_end = date(year, month, w_end_day)

            w_in = 0
            w_out = 0
            for item in items:
                if item.data_inicio and w_start <= item.data_inicio <= w_end:
                    w_in += 1
                if item.completed_at and w_start <= item.completed_at.date() <= w_end:
                    w_out += 1

            week_stats.append(
                MonthlyWeekStatDTO(
                    label=f"Sem {week_idx} ({w_start.day:02d}-{w_end.day:02d})",
                    data_inicio=w_start,
                    data_fim=w_end,
                    entradas=w_in,
                    saidas=w_out,
                )
            )
            cur_day = w_end_day + 1
            week_idx += 1

        return MonthlyReportSummaryDTO(
            ano=year,
            mes=month,
            nome_mes=MONTH_NAMES[month],
            periodo_inicio=periodo_inicio,
            periodo_fim=periodo_fim,
            data_referencia=today,
            is_mes_fechado=is_mes_fechado,
            total_criadas=created_count,
            total_criadas_pecas=created_pcs,
            total_concluidas=concluded_count,
            total_concluidas_pecas=concluded_pcs,
            concluidas_no_prazo=on_time_count,
            concluidas_com_atraso=delayed_count,
            taxa_pontualidade=taxa_pontualidade,
            lead_time_medio_dias=lead_time_medio,
            em_producao_agora=in_production_now,
            em_atraso_agora=in_delay_now,
            previsao_restante_mes=remaining_forecast,
            setores_stats=tuple(sector_stats),
            semanas_stats=tuple(week_stats),
            ops=tuple(items),
        )

    def export_to_excel(self, summary: MonthlyReportSummaryDTO, file_path: str | Path) -> Path:
        """Gera uma pasta de trabalho Excel (.xlsx) altamente profissional e diagramada."""
        target = Path(file_path).resolve()
        target.parent.mkdir(parents=True, exist_ok=True)

        wb = openpyxl.Workbook()

        # Estilos reutilizáveis
        header_fill = PatternFill(start_color="1E3A8A", end_color="1E3A8A", fill_type="solid")
        header_font = Font(name="Segoe UI", size=11, bold=True, color="FFFFFF")
        sub_fill = PatternFill(start_color="F1F5F9", end_color="F1F5F9", fill_type="solid")

        thin_side = Side(border_style="thin", color="CBD5E1")
        cell_border = Border(left=thin_side, right=thin_side, top=thin_side, bottom=thin_side)

        align_center = Alignment(horizontal="center", vertical="center")
        align_left = Alignment(horizontal="left", vertical="center")
        align_right = Alignment(horizontal="right", vertical="center")

        # -------------------------------------------------------------
        # ABA 1: RESUMO EXECUTIVO
        # -------------------------------------------------------------
        ws_resumo = wb.active
        ws_resumo.title = "Resumo Executivo"
        ws_resumo.views.sheetView[0].showGridLines = True

        # Título
        ws_resumo.merge_cells("A1:E1")
        title_cell = ws_resumo["A1"]
        title_cell.value = f"PRODUÇÃO OPERACIONAL — RELATÓRIO DE {summary.nome_mes.upper()} / {summary.ano}"
        title_cell.font = Font(name="Segoe UI", size=14, bold=True, color="1E3A8A")
        title_cell.alignment = align_left

        sit_desc = "MÊS FECHADO / CONCLUÍDO" if summary.is_mes_fechado else f"PARCIAL EM ANDAMENTO (até {format_br_date(summary.data_referencia)})"
        ws_resumo["A2"] = f"Situação: {sit_desc}"
        ws_resumo["A2"].font = Font(name="Segoe UI", size=10, italic=True, color="64748B")

        # Bloco de KPIs
        ws_resumo["A4"] = "INDICADOR"
        ws_resumo["B4"] = "VALOR"
        ws_resumo["C4"] = "DETALHES"
        for col_name in ("A4", "B4", "C4"):
            ws_resumo[col_name].fill = header_fill
            ws_resumo[col_name].font = header_font
            ws_resumo[col_name].alignment = align_center

        kpis = [
            ("Total de OPs Criadas no Mês", f"{summary.total_criadas} OPs", f"{summary.total_criadas_pecas} peças"),
            ("Total de OPs Concluídas no Mês", f"{summary.total_concluidas} OPs", f"{summary.total_concluidas_pecas} peças"),
            ("Entregas Rigorosamente no Prazo", f"{summary.concluidas_no_prazo} OPs", f"Pontualidade: {summary.taxa_pontualidade:.1f}%"),
            ("Entregas com Atraso", f"{summary.concluidas_com_atraso} OPs", f"Lead Time Médio: {summary.lead_time_medio_dias:.1f} dias"),
            ("Em Produção no Chão de Fábrica", f"{summary.em_producao_agora} OPs", f"{summary.em_atraso_agora} atualmente em atraso"),
            ("Previsão de Entrega até Fim do Mês", f"{summary.previsao_restante_mes} OPs", "Entregas programadas restantes"),
        ]

        for i, (lbl, val, det) in enumerate(kpis, start=5):
            ws_resumo[f"A{i}"] = lbl
            ws_resumo[f"B{i}"] = val
            ws_resumo[f"C{i}"] = det
            ws_resumo[f"A{i}"].font = Font(name="Segoe UI", size=10, bold=True)
            ws_resumo[f"B{i}"].font = Font(name="Segoe UI", size=10, bold=True, color="1E3A8A")
            ws_resumo[f"B{i}"].alignment = align_center
            ws_resumo[f"C{i}"].font = Font(name="Segoe UI", size=9, color="475569")
            for c in (f"A{i}", f"B{i}", f"C{i}"):
                ws_resumo[c].border = cell_border
                if i % 2 == 1:
                    ws_resumo[c].fill = sub_fill

        # Tabela por Setor
        start_sec = 13
        ws_resumo[f"A{start_sec}"] = "SETOR"
        ws_resumo[f"B{start_sec}"] = "TOTAL OPS"
        ws_resumo[f"C{start_sec}"] = "TOTAL PEÇAS"
        for col_name in (f"A{start_sec}", f"B{start_sec}", f"C{start_sec}"):
            ws_resumo[col_name].fill = header_fill
            ws_resumo[col_name].font = header_font
            ws_resumo[col_name].alignment = align_center

        cur_row = start_sec + 1
        for sec in summary.setores_stats:
            ws_resumo[f"A{cur_row}"] = sec.setor_nome
            ws_resumo[f"B{cur_row}"] = sec.total_ops
            ws_resumo[f"C{cur_row}"] = sec.total_quantidade
            ws_resumo[f"A{cur_row}"].font = Font(name="Segoe UI", size=10)
            ws_resumo[f"B{cur_row}"].font = Font(name="Segoe UI", size=10, bold=True)
            ws_resumo[f"B{cur_row}"].alignment = align_center
            ws_resumo[f"C{cur_row}"].font = Font(name="Segoe UI", size=10)
            ws_resumo[f"C{cur_row}"].alignment = align_center
            for c in (f"A{cur_row}", f"B{cur_row}", f"C{cur_row}"):
                ws_resumo[c].border = cell_border
            cur_row += 1

        # Auto-ajuste de colunas na aba de resumo
        for col in ws_resumo.columns:
            max_len = max(len(str(cell.value or "")) for cell in col)
            col_letter = get_column_letter(col[0].column)
            ws_resumo.column_dimensions[col_letter].width = max(max_len + 4, 16)

        # -------------------------------------------------------------
        # ABA 2: TODAS AS ORDENS DE PRODUÇÃO (DETALHADO)
        # -------------------------------------------------------------
        ws_ops = wb.create_sheet(title="Ordens de Produção")
        ws_ops.views.sheetView[0].showGridLines = True

        headers = [
            "OP",
            "Cliente",
            "Modelo",
            "Quantidade",
            "Voltagem",
            "Setor",
            "Data Início",
            "Prazo Entrega",
            "Data Conclusão",
            "Ciclo (Dias)",
            "Situação",
        ]

        ws_ops.append(headers)
        for col_idx in range(1, len(headers) + 1):
            cell = ws_ops.cell(row=1, column=col_idx)
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = align_center
            cell.border = cell_border

        for op in summary.ops:
            dt_ini = format_br_date(op.data_inicio) or "-"
            dt_ent = format_br_date(op.data_entrega) or "-"
            dt_conc = op.completed_at.strftime("%d/%m/%Y") if op.completed_at else "-"
            ciclo = op.dias_producao if op.dias_producao is not None else "-"

            sit = "Em Linha"
            if op.categoria == "CONCLUIDA_NO_PRAZO":
                sit = "Concluída no Prazo"
            elif op.categoria == "CONCLUIDA_COM_ATRASO":
                sit = "Concluída com Atraso"
            elif op.categoria == "EM_ATRASO":
                sit = "Em Linha (Atrasada)"

            row_data = [
                op.numero_op,
                op.cliente,
                op.modelo,
                op.quantidade if op.quantidade is not None else "-",
                op.voltagem,
                op.setor_nome,
                dt_ini,
                dt_ent,
                dt_conc,
                ciclo,
                sit,
            ]
            ws_ops.append(row_data)

            r_idx = ws_ops.max_row
            for col_idx in range(1, len(headers) + 1):
                c = ws_ops.cell(row=r_idx, column=col_idx)
                c.border = cell_border
                c.font = Font(name="Segoe UI", size=10)

                # Alinhamentos
                if col_idx in {1, 4, 7, 8, 9, 10, 11}:
                    c.alignment = align_center

                # Cores de situação
                if col_idx == 11:
                    if op.categoria == "CONCLUIDA_NO_PRAZO":
                        c.fill = PatternFill(start_color="DCFCE7", end_color="DCFCE7", fill_type="solid")
                        c.font = Font(name="Segoe UI", size=10, bold=True, color="15803D")
                    elif op.categoria == "CONCLUIDA_COM_ATRASO":
                        c.fill = PatternFill(start_color="FEE2E2", end_color="FEE2E2", fill_type="solid")
                        c.font = Font(name="Segoe UI", size=10, bold=True, color="B91C1C")
                    elif op.categoria == "EM_ATRASO":
                        c.fill = PatternFill(start_color="FFE4E6", end_color="FFE4E6", fill_type="solid")
                        c.font = Font(name="Segoe UI", size=10, bold=True, color="DC2626")
                    else:
                        c.fill = PatternFill(start_color="DBEAFE", end_color="DBEAFE", fill_type="solid")
                        c.font = Font(name="Segoe UI", size=10, bold=True, color="1D4ED8")

        # Auto-ajuste de colunas garantindo que NUNCA apareça "###"
        for col in ws_ops.columns:
            max_len = 0
            for cell in col:
                val_str = str(cell.value or "")
                if len(val_str) > max_len:
                    max_len = len(val_str)
            col_letter = get_column_letter(col[0].column)
            ws_ops.column_dimensions[col_letter].width = max(max_len + 4, 13)

        # Congela cabeçalho e ativa autofiltro
        ws_ops.freeze_panes = "A2"
        ws_ops.auto_filter.ref = ws_ops.dimensions

        wb.save(target)
        return target

    def export_to_csv(self, summary: MonthlyReportSummaryDTO, file_path: str | Path) -> Path:
        """Compatibilidade: exporta dados em CSV."""
        target = Path(file_path).resolve()
        target.parent.mkdir(parents=True, exist_ok=True)

        with target.open("w", newline="", encoding="utf-8-sig") as f:
            writer = csv.writer(f, delimiter=";")
            writer.writerow(["RELATÓRIO MENSAL DE PRODUÇÃO", f"{summary.nome_mes} / {summary.ano}"])
            status_desc = "MÊS FECHADO" if summary.is_mes_fechado else f"PARCIAL EM ANDAMENTO (até {format_br_date(summary.data_referencia)})"
            writer.writerow(["Situação", status_desc])
            writer.writerow([])
            writer.writerow(["INDICADOR", "VALOR"])
            writer.writerow(["Total de OPs Criadas no Mês", summary.total_criadas])
            writer.writerow(["Volume de Peças Criadas", summary.total_criadas_pecas])
            writer.writerow(["Total de OPs Concluídas no Mês", summary.total_concluidas])
            writer.writerow(["Volume de Peças Concluídas", summary.total_concluidas_pecas])
            writer.writerow(["Concluídas no Prazo", summary.concluidas_no_prazo])
            writer.writerow(["Concluídas com Atraso", summary.concluidas_com_atraso])
            writer.writerow(["Índice de Pontualidade (%)", f"{summary.taxa_pontualidade}%"])
            writer.writerow(["Lead Time Médio (Dias)", f"{summary.lead_time_medio_dias} dias"])
            writer.writerow([])

            writer.writerow([
                "OP",
                "Cliente",
                "Modelo",
                "Quantidade",
                "Voltagem",
                "Setor",
                "Data Início",
                "Prazo Entrega",
                "Data Conclusão",
                "Dias Produção",
                "Situação",
            ])

            for op in summary.ops:
                dt_conc = op.completed_at.strftime("%d/%m/%Y") if op.completed_at else ""
                sit = "Em Linha"
                if op.categoria == "CONCLUIDA_NO_PRAZO":
                    sit = "Concluída no Prazo"
                elif op.categoria == "CONCLUIDA_COM_ATRASO":
                    sit = "Concluída com Atraso"
                elif op.categoria == "EM_ATRASO":
                    sit = "Em Linha (Atrasada)"

                writer.writerow([
                    op.numero_op,
                    op.cliente,
                    op.modelo,
                    op.quantidade if op.quantidade is not None else "",
                    op.voltagem,
                    op.setor_nome,
                    format_br_date(op.data_inicio),
                    format_br_date(op.data_entrega),
                    dt_conc,
                    op.dias_producao if op.dias_producao is not None else "",
                    sit,
                ])

        return target

    def export_to_pdf(self, summary: MonthlyReportSummaryDTO, file_path: str | Path) -> Path:
        """Gera um PDF executivo de alta fidelidade visual usando o motor Chromium do Windows ou Qt."""
        target = Path(file_path).resolve()
        target.parent.mkdir(parents=True, exist_ok=True)

        html_content = self._generate_html_report(summary)

        # Salva arquivo HTML temporário
        temp_html = target.with_suffix(".tmp.html")
        temp_html.write_text(html_content, encoding="utf-8")

        # 1. Tenta usar o Microsoft Edge headless nativo do Windows (qualidade vetorial perfeita 100% da folha A4)
        edge_paths = [
            Path(r"C:\Program Files (x86)\Microsoft\Edge\Application\msedge.exe"),
            Path(r"C:\Program Files\Microsoft\Edge\Application\msedge.exe"),
        ]
        edge_bin = next((p for p in edge_paths if p.is_file()), None)

        if edge_bin:
            try:
                in_uri = f"file:///{str(temp_html).replace('\\', '/')}"
                cmd = [
                    str(edge_bin),
                    "--headless",
                    "--disable-gpu",
                    "--no-pdf-header-footer",
                    f"--print-to-pdf={str(target)}",
                    in_uri,
                ]
                res = subprocess.run(cmd, capture_output=True, timeout=15)
                if target.is_file() and target.stat().st_size > 1000:
                    temp_html.unlink(missing_ok=True)
                    return target
            except Exception:
                pass

        # 2. Fallback: PySide6 QPdfWriter calibrado com 96 DPI para ocupar 100% da folha A4
        try:
            from PySide6.QtCore import QMarginsF, QSizeF
            from PySide6.QtGui import QFont, QPageLayout, QPageSize, QPdfWriter, QTextDocument

            doc = QTextDocument()
            doc.setDefaultFont(QFont("Segoe UI", 9))
            doc.setHtml(html_content)

            writer = QPdfWriter(str(target))
            writer.setResolution(96)  # Calibra para escala de tela 96 DPI
            writer.setPageSize(QPageSize(QPageSize.PageSizeId.A4))
            writer.setPageOrientation(QPageLayout.Orientation.Landscape)
            writer.setPageMargins(QMarginsF(10, 10, 10, 10), QPageLayout.Unit.Millimeter)

            # Ajusta tamanho exato do documento para preencher o canvas de impressão
            doc.setPageSize(QSizeF(writer.width(), writer.height()))
            doc.print_(writer)
            del writer
        finally:
            temp_html.unlink(missing_ok=True)

        return target

    def _generate_html_report(self, s: MonthlyReportSummaryDTO) -> str:
        status_badge = (
            "<span style='background: #14532d; color: #86efac; border: 1px solid #166534; padding: 6px 14px; border-radius: 6px; font-weight: bold; font-size: 11px;'>MÊS FECHADO / CONCLUÍDO</span>"
            if s.is_mes_fechado
            else f"<span style='background: #78350f; color: #fde047; border: 1px solid #92400e; padding: 6px 14px; border-radius: 6px; font-weight: bold; font-size: 11px;'>PARCIAL EM ANDAMENTO (até {format_br_date(s.data_referencia)})</span>"
        )

        rows_html = []
        for idx, op in enumerate(s.ops):
            dt_conc = op.completed_at.strftime("%d/%m/%Y") if op.completed_at else "-"
            bg_tr = "#ffffff" if idx % 2 == 0 else "#f8fafc"

            if op.categoria == "CONCLUIDA_NO_PRAZO":
                badge = "<span style='background: #dcfce7; color: #15803d; border: 1px solid #bbf7d0; padding: 3px 8px; border-radius: 12px; font-weight: bold; font-size: 10px;'>✔ No Prazo</span>"
            elif op.categoria == "CONCLUIDA_COM_ATRASO":
                badge = "<span style='background: #fee2e2; color: #b91c1c; border: 1px solid #fecaca; padding: 3px 8px; border-radius: 12px; font-weight: bold; font-size: 10px;'>✖ Com Atraso</span>"
            elif op.categoria == "EM_ATRASO":
                badge = "<span style='background: #ffe4e6; color: #dc2626; border: 1px solid #fecdd3; padding: 3px 8px; border-radius: 12px; font-weight: bold; font-size: 10px;'>🚨 Atrasada Hoje</span>"
            else:
                badge = "<span style='background: #dbeafe; color: #1d4ed8; border: 1px solid #bfdbfe; padding: 3px 8px; border-radius: 12px; font-weight: bold; font-size: 10px;'>⚙ Em Linha</span>"

            rows_html.append(
                f"""
                <tr style='background-color: {bg_tr};'>
                    <td style='font-weight: 700; padding: 7px 8px; border-bottom: 1px solid #e2e8f0; text-align: center; color: #0f172a;'>{op.numero_op}</td>
                    <td style='padding: 7px 8px; border-bottom: 1px solid #e2e8f0; color: #1e293b; font-weight: 600;'>{op.cliente}</td>
                    <td style='padding: 7px 8px; border-bottom: 1px solid #e2e8f0; color: #334155;'>{op.modelo}</td>
                    <td style='padding: 7px 8px; border-bottom: 1px solid #e2e8f0; text-align: center; font-weight: bold; color: #0f172a;'>{op.quantidade or '-'}</td>
                    <td style='padding: 7px 8px; border-bottom: 1px solid #e2e8f0; color: #475569;'>{op.setor_nome}</td>
                    <td style='padding: 7px 8px; border-bottom: 1px solid #e2e8f0; text-align: center; color: #64748b;'>{format_br_date(op.data_inicio) or '-'}</td>
                    <td style='padding: 7px 8px; border-bottom: 1px solid #e2e8f0; text-align: center; color: #64748b;'>{format_br_date(op.data_entrega) or '-'}</td>
                    <td style='padding: 7px 8px; border-bottom: 1px solid #e2e8f0; text-align: center; color: #64748b;'>{dt_conc}</td>
                    <td style='padding: 7px 8px; border-bottom: 1px solid #e2e8f0; text-align: center; font-weight: 600; color: #0f172a;'>{op.dias_producao if op.dias_producao is not None else '-'} d</td>
                    <td style='padding: 7px 8px; border-bottom: 1px solid #e2e8f0; text-align: center;'>{badge}</td>
                </tr>
                """
            )

        # Gráfico Donut em SVG para o PDF
        total_conc = s.concluidas_no_prazo + s.concluidas_com_atraso
        r = 50
        cx, cy = 65, 65
        circ = 2 * 3.14159 * r
        stroke_dash = (s.taxa_pontualidade / 100.0) * circ
        donut_svg = f"""
        <svg width="130" height="130" viewBox="0 0 130 130">
            <circle cx="{cx}" cy="{cy}" r="{r}" fill="none" stroke="#ef4444" stroke-width="16" />
            <circle cx="{cx}" cy="{cy}" r="{r}" fill="none" stroke="#22c55e" stroke-width="16"
                    stroke-dasharray="{stroke_dash} {circ}" stroke-dashoffset="{circ * 0.25}" />
            <text x="{cx}" y="{cy - 4}" text-anchor="middle" font-size="18" font-weight="bold" fill="#0f172a" font-family="Segoe UI">{s.taxa_pontualidade:.1f}%</text>
            <text x="{cx}" y="{cy + 16}" text-anchor="middle" font-size="10" font-weight="600" fill="#64748b" font-family="Segoe UI">No Prazo</text>
        </svg>
        """

        # Barras de semanas
        weeks_html = []
        max_ops = max([max(w.entradas, w.saidas) for w in s.semanas_stats] + [1])
        for w in s.semanas_stats:
            in_pct = int((w.entradas / max_ops) * 100)
            out_pct = int((w.saidas / max_ops) * 100)
            weeks_html.append(
                f"""
                <div style='margin-bottom: 7px;'>
                    <div style='font-size: 11px; font-weight: 700; color: #334155; margin-bottom: 2px;'>{w.label} &bull; {w.entradas} entradas / {w.saidas} saídas</div>
                    <div style='background: #e2e8f0; height: 12px; border-radius: 4px; overflow: hidden; display: flex;'>
                        <div style='background: #3b82f6; width: {in_pct}%; height: 12px; text-align: center; color: white; font-size: 9px; line-height: 12px;'></div>
                        <div style='background: #10b981; width: {out_pct}%; height: 12px; text-align: center; color: white; font-size: 9px; line-height: 12px;'></div>
                    </div>
                </div>
                """
            )

        # Setores
        sector_bars = []
        max_sec = max([sec.total_ops for sec in s.setores_stats] + [1])
        for sec in s.setores_stats:
            pct = int((sec.total_ops / max_sec) * 100)
            sector_bars.append(
                f"""
                <div style='margin-bottom: 6px;'>
                    <div style='font-size: 11px; color: #1e293b; margin-bottom: 2px;'><strong>{sec.setor_nome}:</strong> {sec.total_ops} OPs ({sec.total_quantidade} peças)</div>
                    <div style='background: #e2e8f0; height: 10px; border-radius: 4px;'>
                        <div style='background: {sec.cor}; width: {pct}%; height: 10px; border-radius: 4px;'></div>
                    </div>
                </div>
                """
            )

        html = f"""<!DOCTYPE html>
        <html>
        <head>
            <meta charset="utf-8">
            <title>Relatório Mensal - {s.nome_mes} / {s.ano}</title>
            <style>
                @page {{
                    size: A4 landscape;
                    margin: 10mm 12mm 10mm 12mm;
                }}
                * {{ box-sizing: border-box; }}
                body {{
                    font-family: 'Segoe UI', Arial, sans-serif;
                    color: #0f172a;
                    background: #ffffff;
                    margin: 0;
                    padding: 0;
                    -webkit-print-color-adjust: exact;
                    print-color-adjust: exact;
                }}
                .header-table {{ width: 100%; border-bottom: 2px solid #2563eb; padding-bottom: 12px; margin-bottom: 14px; }}
                .kpi-row {{ display: table; width: 100%; margin-bottom: 14px; table-layout: fixed; }}
                .kpi-cell {{ display: table-cell; padding: 0 4px; }}
                .kpi-card {{
                    background: #f8fafc; border: 1px solid #cbd5e1; border-radius: 6px;
                    padding: 10px 12px; text-align: left;
                }}
                .kpi-title {{ font-size: 9px; font-weight: 800; color: #64748b; text-transform: uppercase; letter-spacing: 0.5px; }}
                .kpi-num {{ font-size: 20px; font-weight: 900; color: #0f172a; margin: 3px 0; }}
                .kpi-sub {{ font-size: 10px; font-weight: 600; color: #475569; }}

                .charts-table {{ width: 100%; margin-bottom: 14px; table-layout: fixed; }}
                .chart-panel {{
                    background: #f8fafc; border: 1px solid #cbd5e1; border-radius: 8px;
                    padding: 12px; height: 100%;
                }}
                .panel-title {{ font-size: 12px; font-weight: 800; color: #0f172a; margin-bottom: 8px; border-bottom: 1px solid #e2e8f0; padding-bottom: 4px; }}

                table.data {{ width: 100%; border-collapse: collapse; font-size: 10px; page-break-inside: auto; }}
                table.data th {{
                    background-color: #1e3a8a; color: #ffffff; font-weight: 700;
                    padding: 8px 6px; text-align: left;
                }}
                table.data tr {{ page-break-inside: avoid; page-break-after: auto; }}
            </style>
        </head>
        <body>
            <!-- Cabeçalho -->
            <table class='header-table'>
                <tr>
                    <td style='vertical-align: middle;'>
                        <div style='font-size: 22px; font-weight: 900; color: #1e3a8a; letter-spacing: 0.5px;'>PRODUÇÃO OPERACIONAL</div>
                        <div style='font-size: 12px; color: #64748b; font-weight: 600;'>Relatório Executivo de Desempenho e Indicadores Industriais</div>
                    </td>
                    <td style='vertical-align: middle; text-align: right;'>
                        <div style='font-size: 20px; font-weight: 900; color: #0f172a;'>{s.nome_mes.upper()} / {s.ano}</div>
                        <div style='margin-top: 5px;'>{status_badge}</div>
                    </td>
                </tr>
            </table>

            <!-- 5 KPIs em Grid -->
            <div class='kpi-row'>
                <div class='kpi-cell'>
                    <div class='kpi-card' style='border-left: 4px solid #3b82f6;'>
                        <div class='kpi-title'>Entradas no Mês</div>
                        <div class='kpi-num'>{s.total_criadas} OPs</div>
                        <div class='kpi-sub'>{s.total_criadas_pecas} peças</div>
                    </div>
                </div>
                <div class='kpi-cell'>
                    <div class='kpi-card' style='border-left: 4px solid #10b981;'>
                        <div class='kpi-title'>Concluídas</div>
                        <div class='kpi-num'>{s.total_concluidas} OPs</div>
                        <div class='kpi-sub'>{s.total_concluidas_pecas} peças</div>
                    </div>
                </div>
                <div class='kpi-cell'>
                    <div class='kpi-card' style='border-left: 4px solid #22c55e;'>
                        <div class='kpi-title'>No Prazo</div>
                        <div class='kpi-num' style='color: #15803d;'>{s.concluidas_no_prazo} OPs</div>
                        <div class='kpi-sub' style='color: #15803d;'>Pontualidade: {s.taxa_pontualidade:.1f}%</div>
                    </div>
                </div>
                <div class='kpi-cell'>
                    <div class='kpi-card' style='border-left: 4px solid #ef4444;'>
                        <div class='kpi-title'>Com Atraso</div>
                        <div class='kpi-num' style='color: #b91c1c;'>{s.concluidas_com_atraso} OPs</div>
                        <div class='kpi-sub' style='color: #b91c1c;'>Lead Time: {s.lead_time_medio_dias:.1f}d</div>
                    </div>
                </div>
                <div class='kpi-cell'>
                    <div class='kpi-card' style='border-left: 4px solid #f59e0b;'>
                        <div class='kpi-title'>Em Produção Hoje</div>
                        <div class='kpi-num' style='color: #b45309;'>{s.em_producao_agora} OPs</div>
                        <div class='kpi-sub' style='color: #dc2626;'>{s.em_atraso_agora} em atraso</div>
                    </div>
                </div>
            </div>

            <!-- Painéis com Gráficos -->
            <table class='charts-table'>
                <tr>
                    <td style='width: 25%; vertical-align: top; padding-right: 8px;'>
                        <div class='chart-panel' style='text-align: center;'>
                            <div class='panel-title' style='text-align: left;'>🎯 Pontualidade</div>
                            <div style='margin: 4px 0;'>{donut_svg}</div>
                            <div style='font-size: 10px; color: #475569; font-weight: bold;'>
                                <span style='color: #22c55e;'>✔ {s.concluidas_no_prazo} no prazo</span> &bull; <span style='color: #ef4444;'>✖ {s.concluidas_com_atraso} atrasadas</span>
                            </div>
                        </div>
                    </td>
                    <td style='width: 40%; vertical-align: top; padding-right: 8px;'>
                        <div class='chart-panel'>
                            <div class='panel-title'>📊 Fluxo Semanal (Entradas vs Conclusões)</div>
                            <div style='font-size: 10px; margin-bottom: 6px; color: #64748b;'>
                                <span style='color: #3b82f6;'>■ Entradas</span> &bull; <span style='color: #10b981;'>■ Conclusões</span>
                            </div>
                            {''.join(weeks_html) if weeks_html else '<em>Sem movimentação semanal</em>'}
                        </div>
                    </td>
                    <td style='width: 35%; vertical-align: top;'>
                        <div class='chart-panel'>
                            <div class='panel-title'>🏭 Distribuição por Setores</div>
                            <div style='margin-top: 6px;'>
                                {''.join(sector_bars) if sector_bars else '<em>Sem OPs alocadas</em>'}
                            </div>
                        </div>
                    </td>
                </tr>
            </table>

            <!-- Tabela Detalhada -->
            <div style='margin-top: 10px;'>
                <div style='font-size: 12px; font-weight: 800; color: #0f172a; margin-bottom: 6px;'>
                    Detalhamento das Ordens de Produção ({len(s.ops)} registros)
                </div>
                <table class='data'>
                    <thead>
                        <tr>
                            <th style='width: 60px; text-align: center;'>OP</th>
                            <th>Cliente</th>
                            <th>Modelo</th>
                            <th style='width: 45px; text-align: center;'>Qtd</th>
                            <th>Setor</th>
                            <th style='width: 75px; text-align: center;'>Início</th>
                            <th style='width: 75px; text-align: center;'>Entrega</th>
                            <th style='width: 75px; text-align: center;'>Conclusão</th>
                            <th style='width: 50px; text-align: center;'>Ciclo</th>
                            <th style='width: 100px; text-align: center;'>Situação</th>
                        </tr>
                    </thead>
                    <tbody>
                        {''.join(rows_html) if rows_html else '<tr><td colspan=\"10\" style=\"text-align: center; padding: 15px; color: #64748b;\">Nenhuma OP encontrada no período.</td></tr>'}
                    </tbody>
                </table>
            </div>

            <!-- Rodapé -->
            <div style='margin-top: 18px; border-top: 1px solid #cbd5e1; padding-top: 6px; font-size: 9px; color: #94a3b8; display: table; width: 100%;'>
                <div style='display: table-cell;'>Produção Operacional &bull; Sistema de Gestão Industrial</div>
                <div style='display: table-cell; text-align: right;'>Documento gerado em {datetime.now().strftime('%d/%m/%Y às %H:%M')}</div>
            </div>
        </body>
        </html>
        """
        return html
